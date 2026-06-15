#!/usr/bin/env python3
"""Minimal Markdown -> DOCX converter for the Technical Design Review.

Handles: ATX headings, blockquotes, bullet lists, fenced code blocks
(including mermaid, rendered as monospaced text), pipe tables, and inline
**bold** / `code`. Not a general Markdown engine — just enough for our doc.
"""
from __future__ import annotations

import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor


def add_inline(paragraph, text: str) -> None:
    """Render **bold** and `code` spans into runs."""
    for token in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(token)


def cell_text(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text.replace("\\|", "|").strip()


def convert(md_path: str, docx_path: str) -> None:
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]

        # fenced code block (``` ... ```), incl. mermaid
        if line.lstrip().startswith("```"):
            lang = line.lstrip()[3:].strip()
            i += 1
            code: list[str] = []
            while i < n and not lines[i].lstrip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            if lang:
                cap = doc.add_paragraph()
                r = cap.add_run(f"[{lang} diagram]" if lang == "mermaid" else f"[{lang}]")
                r.italic = True
                r.font.size = Pt(8)
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            continue

        # pipe table
        if line.strip().startswith("|") and i + 1 < n and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            header = [cell_text(c) for c in line.strip().strip("|").split("|")]
            i += 2  # skip header + separator
            rows: list[list[str]] = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append([cell_text(c) for c in lines[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Light Grid Accent 1"
            for j, h in enumerate(header):
                c = t.rows[0].cells[j]
                c.text = h
                for para in c.paragraphs:
                    for run in para.runs:
                        run.bold = True
                        run.font.size = Pt(9)
            for row in rows:
                cells = t.add_row().cells
                for j in range(len(header)):
                    cells[j].text = row[j] if j < len(row) else ""
                    for para in cells[j].paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(9)
            doc.add_paragraph()
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1)) - 1  # # -> Title(0), ## -> 1, ...
            doc.add_heading(cell_text(m.group(2)), level=min(level, 4))
            i += 1
            continue

        # horizontal rule
        if re.match(r"^\s*---+\s*$", line):
            i += 1
            continue

        # blockquote
        if line.startswith(">"):
            p = doc.add_paragraph()
            r = p.add_run(line.lstrip("> ").rstrip())
            r.italic = True
            i += 1
            continue

        # bullet
        if re.match(r"^\s*[-*]\s+", line):
            text = re.sub(r"^\s*[-*]\s+", "", line)
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, text)
            i += 1
            continue

        # numbered list
        if re.match(r"^\s*\d+\.\s+", line):
            text = re.sub(r"^\s*\d+\.\s+", "", line)
            p = doc.add_paragraph(style="List Number")
            add_inline(p, text)
            i += 1
            continue

        # blank
        if not line.strip():
            i += 1
            continue

        # paragraph
        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1

    doc.save(docx_path)
    print("WROTE:", docx_path)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
